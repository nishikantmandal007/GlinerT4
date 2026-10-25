# Experiment log for the GLiNER2 T4 kernel work: one row per measurement in kernel_log.xlsx.
#   uv run --with openpyxl log.py '<json row or list of rows>'
# Row keys: phase, experiment, change, config, metric, value, unit, baseline, accuracy, outcome, notes.
import datetime
import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

PATH = Path(__file__).with_name("kernel_log.xlsx")
COLUMNS = ["id", "time", "phase", "experiment", "change", "config", "metric", "value", "unit", "baseline", "speedup",
           "accuracy", "outcome", "notes"]
WIDTHS = [5, 17, 14, 34, 50, 30, 22, 10, 8, 10, 9, 40, 12, 60]
OUTCOME_FILL = {"kept": "C6EFCE", "rejected": "FFC7CE", "baseline": "DDEBF7", "info": "F2F2F2"}


def open_sheet():
    if PATH.exists():
        book = load_workbook(PATH)
        return book, book.active
    book = Workbook()
    sheet = book.active
    sheet.title = "experiments"
    sheet.append(COLUMNS)
    for cell, width in zip(sheet[1], WIDTHS):
        cell.font = Font(bold=True)
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.freeze_panes = "A2"
    return book, sheet


def add(rows):
    book, sheet = open_sheet()
    for row in rows:
        value, baseline = row.get("value"), row.get("baseline")
        # lower is better for times, higher for throughput
        speedup = None
        if isinstance(value, (int, float)) and isinstance(baseline, (int, float)) and value and baseline:
            speedup = round(baseline / value if row.get("unit") in ("ms", "s") else value / baseline, 3)
        sheet.append([sheet.max_row, datetime.datetime.now().strftime("%Y-%m-%d %H:%M"), *(row.get(key) for key in COLUMNS[2:10]),
                      speedup, *(row.get(key) for key in COLUMNS[11:])])
        for cell in sheet[sheet.max_row]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        fill = OUTCOME_FILL.get(row.get("outcome"))
        if fill:
            sheet.cell(sheet.max_row, COLUMNS.index("outcome") + 1).fill = PatternFill("solid", fgColor=fill)
    book.save(PATH)
    print(f"{len(rows)} row(s) -> {PATH} ({sheet.max_row - 1} total)")


if __name__ == "__main__":
    data = json.loads(sys.argv[1])
    add(data if isinstance(data, list) else [data])
