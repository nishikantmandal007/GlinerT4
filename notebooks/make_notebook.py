# Writes notebooks/glinert4_speed.ipynb (no outputs). Run it in Colab on a T4 to fill the outputs.
#   python notebooks/make_notebook.py
import json
from pathlib import Path

REPO = "https://github.com/nishikantmandal007/GlinerT4"
REF = "claude/upbeat-bell-km1ebk"

cells = []


def md(text):
    cells.append({"cell_type": "markdown", "metadata": {}, "source": text.strip("\n")})


def code(text):
    cells.append({"cell_type": "code", "metadata": {}, "execution_count": None, "outputs": [],
                  "source": text.strip("\n")})


md(f"""
# Making GLiNER2 fast on a T4 without losing accuracy

I wanted GLiNER2 to run fast on a Tesla T4, the cheapest GPU on most clouds, without changing a single prediction.
This notebook runs the same model on the same text four times, adding one optimisation at a time, and scores every
run against the gold labels. The starting point is GLiNER2 in fp16: in my tests fp32 was slower and no more accurate.

| Step | What changes |
|---|---|
| fp16 | GLiNER2 with half-precision weights (`quantize=True`) |
| + disentangled_flash | GLiNER2's fused Triton attention (`attention_backend="disentangled_flash"`) |
| + glinert4 kernels | my CUDA attention kernel for the T4, residual add fused into LayerNorm, a count head that skips work it throws away, a word-split cache |
| + pipeline | the CPU prepares batch *n+1* while the GPU runs batch *n*, and lengths are padded to multiples of 32 |

The text is [CrossNER](https://github.com/zliucr/CrossNER), the zero-shot NER benchmark from the GLiNER2 paper: its
five English test sets, once as single sentences and once joined into documents of about 1,000 words. Every label
of a domain goes in one prompt, threshold 0.5, strict F1 (label, start and end must match).

Pinned: gliner2 2.1.0, `fastino/gliner2-multi-v1`, disentangled-flash 1.1.1, glinert4 from [{REPO}]({REPO}).
Section 5 runs both on a whole document of your own (.docx, .pdf or .txt) and, if you have it, scores them against
your ground truth page by page. Pick a GPU runtime (Runtime → Change runtime type → T4). It takes about 15 minutes.
""")

code(f"""
import sys

if "google.colab" in sys.modules:
    !pip install -q uv
    !uv pip install --system -q "gliner2[local]==2.1.0" disentangled-flash==1.1.1 ninja python-docx pypdf
    !uv pip install --system -q --no-deps "glinert4 @ git+{REPO}@{REF}#subdirectory=production"
""")

code("""
import html
import time
import urllib.request

import ipywidgets as widgets
import matplotlib.pyplot as plt
import torch
import transformers
from IPython.display import HTML, Markdown, display

import gliner2
import glinert4
from gliner2 import AutoExtractor
from glinert4 import pipeline, runtime
from glinert4.ops import ops

CHECKPOINT = "fastino/gliner2-multi-v1"
BATCH, THRESHOLD, DOC_WORDS = 16, 0.5, 1000
CROSSNER = "https://raw.githubusercontent.com/zliucr/CrossNER/main/ner_data/{}/test.txt"
DOMAINS = ("ai", "literature", "music", "politics", "science")
LABEL_NAMES = {"programlang": "programming language", "misc": "miscellaneous", "musicalartist": "musical artist",
               "musicalinstrument": "musical instrument", "musicgenre": "music genre", "literarygenre": "literary genre",
               "politicalparty": "political party", "academicjournal": "academic journal",
               "astronomicalobject": "astronomical object", "chemicalcompound": "chemical compound",
               "chemicalelement": "chemical element"}

ops()
print(f"gliner2 {gliner2.__version__} | transformers {transformers.__version__} | torch {torch.__version__} | "
      f"{torch.cuda.get_device_name()}")


# a markdown table
def table(header, rows):
    lines = ["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
    return "\\n".join(lines + ["| " + " | ".join(str(cell) for cell in row) + " |" for row in rows])
""")

md("""
## 1. Data

CrossNER's test sets straight from GitHub. Documents are consecutive sentences joined until they reach about 1,000
words, with the gold spans moved along, so long-document handling (GLiNER2's 384-word windows with 64 words of
overlap) is scored too.
""")

code("""
# one CrossNER test set as (text, gold spans), a gold span being (label, start, end)
def crossner(domain):
    sentences, tokens, tags = [], [], []
    for line in urllib.request.urlopen(CROSSNER.format(domain)).read().decode().splitlines() + [""]:
        parts = line.split("\\t")
        if len(parts) == 2:
            tokens.append(parts[0])
            tags.append(parts[1])
        elif tokens:
            sentences.append(from_bio(tokens, tags))
            tokens, tags = [], []
    return sentences


# the tokens joined by spaces, and the character spans of their BIO tags
def from_bio(tokens, tags):
    starts = [sum(len(token) + 1 for token in tokens[:i]) for i in range(len(tokens))]
    spans, current = [], None
    for i, tag in enumerate(tags + ["O"]):
        kind = tag[2:] if tag != "O" else None
        if current and (tag.startswith("B-") or kind != current[0]):
            spans.append((LABEL_NAMES.get(current[0], current[0]), starts[current[1]], starts[i - 1] + len(tokens[i - 1])))
            current = None
        if kind and current is None:
            current = (kind, i)
    return " ".join(tokens), spans


# consecutive sentences joined into documents of about DOC_WORDS words
def documents(sentences):
    docs, text, spans = [], "", []
    for sentence, gold in sentences:
        text += " " if text else ""
        spans += [(label, start + len(text), end + len(text)) for label, start, end in gold]
        text += sentence
        if len(text.split()) >= DOC_WORDS:
            docs.append((text, spans))
            text, spans = "", []
    if text:
        docs.append((text, spans))
    return docs


data = {}
for domain in DOMAINS:
    sentences = crossner(domain)
    data[domain] = {"labels": sorted({label for _, gold in sentences for label, _, _ in gold}),
                    "sentences": sentences, "documents": documents(sentences)}

display(Markdown(table(["Domain", "labels", "sentences", "documents", "words"],
                       [[d, len(v["labels"]), len(v["sentences"]), len(v["documents"]),
                         f"{sum(len(t.split()) for t, _ in v['sentences']):,}"] for d, v in data.items()])))
""")

md("""
## 2. Four runs, one optimisation at a time

Every run calls the same GLiNER2 long-text API with the same batch size, labels and threshold. Only the last run
swaps it for glinert4's pipelined version of that API.
""")

code("""
def load(**options):
    return runtime.freeze_eval(AutoExtractor.from_pretrained(CHECKPOINT, map_location="cuda", **options))


def predict(model, texts, labels, pipelined):
    if pipelined:
        return pipeline.extract_entities_long_many(model, [(texts, labels, THRESHOLD)], batch_size=BATCH,
                                                   include_confidence=True, include_spans=True)[0]
    return model.batch_extract_entities_long(texts, labels, batch_size=BATCH, threshold=THRESHOLD,
                                            include_confidence=True, include_spans=True)


def spans(results, labels):
    return [{(label, e["start"], e["end"]) for label in labels for e in result["entities"].get(label, [])}
            for result in results]


# seconds and predicted spans for every text, per kind of input
def measure(model, pipelined=False):
    run = {}
    for kind in ("sentences", "documents"):
        seconds, found = 0.0, []
        for domain in data.values():
            texts = [text for text, _ in domain[kind]]
            predict(model, texts[:BATCH], domain["labels"], pipelined)
            torch.cuda.synchronize()
            start = time.perf_counter()
            found += spans(predict(model, texts, domain["labels"], pipelined), domain["labels"])
            torch.cuda.synchronize()
            seconds += time.perf_counter() - start
        run[kind] = {"seconds": seconds, "found": found}
    return run


runs = {}
stock = load(quantize=True)
runs["fp16"] = measure(stock)
fast = load(quantize=True, attention_backend="disentangled_flash")
runs["+ disentangled_flash"] = measure(fast)
glinert4.install()
runs["+ glinert4 kernels"] = measure(fast)
runtime.add_length_buckets(fast, 32)
runs["+ pipeline"] = measure(fast, pipelined=True)
""")

md("""
## 3. Results

Speed is relative to GLiNER2 fp16. "Same as fp16" is the share of texts whose predicted spans are exactly the ones
GLiNER2 fp16 predicts; anything below 100% is rounding flipping a score that sits right at the threshold.
""")

code("""
def f1(kind, found):
    gold = [set(g) for domain in data.values() for _, g in domain[kind]]
    hits = sum(len(g & p) for g, p in zip(gold, found))
    predicted, expected = sum(map(len, found)), sum(map(len, gold))
    return 2 * hits / (predicted + expected)


for kind in ("sentences", "documents"):
    count = sum(len(domain[kind]) for domain in data.values())
    base = runs["fp16"][kind]
    rows = [[step, f"{run[kind]['seconds']:.1f} s", f"{count / run[kind]['seconds']:.1f}",
             f"**{base['seconds'] / run[kind]['seconds']:.1f}×**", f"{f1(kind, run[kind]['found']):.3f}",
             f"{sum(a == b for a, b in zip(run[kind]['found'], base['found'])) / count:.1%}"]
            for step, run in runs.items()]
    display(Markdown(f"**{count:,} {kind}**\\n\\n" + table(["Step", "time", f"{kind}/s", "speed-up", "F1",
                                                             "same as fp16"], rows)))
""")

code("""
plt.rcParams.update({"font.size": 11, "axes.edgecolor": "#d5d4cf", "axes.labelcolor": "#52514e",
                     "xtick.color": "#52514e", "ytick.color": "#0b0b0b"})
fig, axes = plt.subplots(1, 2, figsize=(12, 3.6), sharex=True, sharey=True)
top = max(runs["fp16"][kind]["seconds"] / run[kind]["seconds"] for run in runs.values() for kind in ("sentences", "documents"))
for ax, kind in zip(axes, ("sentences", "documents")):
    speedup = [runs["fp16"][kind]["seconds"] / run[kind]["seconds"] for run in runs.values()]
    bars = ax.barh(list(runs), speedup, height=0.55, color="#2a78d6")
    ax.bar_label(bars, labels=[f"{s:.1f}×" for s in speedup], padding=5, color="#0b0b0b")
    ax.set_title(f"{kind.capitalize()}: times faster than GLiNER2 fp16", loc="left", color="#0b0b0b", pad=12)
    ax.set_xlim(0, top * 1.15)
    ax.grid(axis="x", color="#ecebe7")
    ax.set_axisbelow(True)
    ax.spines[["top", "right"]].set_visible(False)
axes[0].invert_yaxis()
fig.tight_layout()
plt.show()
""")

md("""
## 4. Try your own text

GLiNER2 fp16 against glinert4 on any text and labels. The timings are the best of three runs.
""")

code("""
PALETTE = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#6250d6", "#e34948"]
EXAMPLE = ("Satya Nadella, the chief executive of Microsoft, met Prime Minister Narendra Modi in New Delhi on "
           "Tuesday to discuss a cloud partnership with Reliance Industries.")


def timed(model, text, labels, pipelined):
    times = []
    for _ in range(3):
        torch.cuda.synchronize()
        start = time.perf_counter()
        found = spans(predict(model, [text], labels, pipelined), labels)[0]
        torch.cuda.synchronize()
        times.append(time.perf_counter() - start)
    return found, min(times) * 1000


def highlight(text, found, labels):
    parts, last = [], 0
    for label, start, end in sorted(found, key=lambda span: span[1]):
        if start < last:
            continue
        color = PALETTE[labels.index(label)]
        parts += [html.escape(text[last:start]),
                  f'<mark style="background:{color}26;border-bottom:2px solid {color};border-radius:3px;'
                  f'padding:1px 3px;color:inherit">{html.escape(text[start:end])}'
                  f'<span style="font-size:0.7em;opacity:0.7;margin-left:4px">{html.escape(label)}</span></mark>']
        last = end
    return "".join(parts + [html.escape(text[last:])])


def compare(text, labels):
    labels = [label.strip() for label in labels.split(",") if label.strip()][:len(PALETTE)]
    glinert4.uninstall()
    before, before_ms = timed(stock, text, labels, False)
    glinert4.install()
    after, after_ms = timed(fast, text, labels, True)
    verdict = "same entities" if before == after else "entities differ"
    rows = [("GLiNER2 fp16", before_ms, before), ("glinert4", after_ms, after)]
    display(HTML("".join(
        f'<div style="margin:10px 0;line-height:1.9"><b>{name}</b> · {ms:.0f} ms<br>{highlight(text, found, labels)}</div>'
        for name, ms, found in rows) + f"<p><b>{before_ms / after_ms:.1f}× faster</b>, {verdict}.</p>"))


compare(EXAMPLE, "person, organization, location, date")
""")

code("""
text_box = widgets.Textarea(value=EXAMPLE, layout=widgets.Layout(width="100%", height="110px"))
label_box = widgets.Text(value="person, organization, location, date", description="Labels",
                         layout=widgets.Layout(width="100%"))
button = widgets.Button(description="Compare", button_style="primary", icon="bolt")
output = widgets.Output()


def on_click(_):
    with output:
        output.clear_output()
        compare(text_box.value, label_box.value)


button.on_click(on_click)
display(widgets.VBox([text_box, label_box, button, output]))
""")

md("""
## 5. Your own document

Both GLiNER2 fp16 and glinert4 run the production detector (Person, Organization, Street, City, one pass, 384-word
windows with 64 words of overlap) on a whole document. A .docx is read page by page (explicit page breaks), with its
tables, tracked deletions, comments, header, footer and document properties, because PII hides in all of them.

Ground truth is optional, in either of two forms:

- a manifest, `{"pii": [{"type": "Names", "value": "…", "location": {"page": 3, "part": "word/document.xml"}}],
  "controls": [{"value": "…"}]}`, scored per page and per occurrence; controls are values that must not be detected;
- plain values, `{"Person": ["…"], "City": ["…"]}`, scored on the whole document.

`Names` counts as Person. Other types (emails, phone numbers, dates) are listed but not scored, because this detector
does not look for them. The default input is a synthetic 100-page service contract with its manifest.
""")

code(r"""
import io
import json
import re
import unicodedata
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from glinert4 import Detector

SAMPLES = "https://raw.githubusercontent.com/nishikantmandal007/GlinerT4/claude/upbeat-bell-km1ebk/notebooks/samples/"
DOCUMENT = "samples/02_service_contract_100p.docx"
GROUND_TRUTH = "samples/02_service_contract_100p.docx.manifest.json"
GT_LABELS = {"Names": "Person", "Person": "Person", "Organization": "Organization", "Street": "Street", "City": "City"}


# a .docx as {page: text} split at explicit page breaks, plus its header, footer, comments and properties
def read_docx(data):
    doc = Document(io.BytesIO(data))
    pages, lines = [], []
    for block in doc.element.body.iterchildren():
        if block.tag == qn("w:p"):
            if block.find(f"{qn('w:pPr')}/{qn('w:pageBreakBefore')}") is not None and lines:
                pages.append("\n".join(lines))
                lines = []
            lines.append(Paragraph(block, doc).text)
            for change in block.iter(qn("w:del")):
                deleted = "".join(t.text or "" for t in change.iter(qn("w:delText")))
                lines.append(f"Deleted by {change.get(qn('w:author'))}: {deleted}")
            if any(br.get(qn("w:type")) == "page" for br in block.iter(qn("w:br"))):
                pages.append("\n".join(lines))
                lines = []
        elif block.tag == qn("w:tbl"):
            lines += ["\t".join(cell.text for cell in row.cells) for row in Table(block, doc).rows]
    if lines:
        pages.append("\n".join(lines))
    parts = dict(enumerate(pages, 1))
    unique = lambda texts: "\n".join(dict.fromkeys(t for t in texts if t.strip()))
    parts["header"] = unique(p.text for s in doc.sections for p in s.header.paragraphs)
    parts["footer"] = unique(p.text for s in doc.sections for p in s.footer.paragraphs)
    parts["comments"] = "\n".join(f"Comment by {c.author}: {c.text}" for c in doc.comments)
    core = doc.core_properties
    parts["properties"] = "\n".join(f"{field.replace('_', ' ').capitalize()}: {getattr(core, field)}"
                                    for field in ("author", "last_modified_by", "title", "subject", "keywords")
                                    if getattr(core, field))
    return {name: text for name, text in parts.items() if text.strip()}


def read_document(name, data):
    if name.lower().endswith(".docx"):
        return read_docx(data)
    if name.lower().endswith(".pdf"):
        from pypdf import PdfReader

        return {page: p.extract_text() or "" for page, p in enumerate(PdfReader(io.BytesIO(data)).pages, 1)}
    return {1: data.decode("utf-8-sig")}


def part_name(location):
    location = location or {}
    part = location.get("part") or "word/document.xml"
    if part.endswith("document.xml"):
        return location.get("page") or 1
    for name in ("header", "footer", "comments"):
        if name in part:
            return name
    return "properties" if part.startswith("docProps/") else part


# Counter{(part, label, value): occurrences}, the control values, and the types this detector does not score
def read_ground_truth(data):
    raw = json.loads(data)
    truth, controls = Counter(), {c["value"] for c in raw.get("controls", [])}
    if "pii" in raw:
        for item in raw["pii"]:
            if item["type"] in GT_LABELS:
                truth[part_name(item.get("location")), GT_LABELS[item["type"]], item["value"]] += 1
        return truth, controls, Counter(item["type"] for item in raw["pii"] if item["type"] not in GT_LABELS)
    for label, values in raw.items():
        for value in values:
            truth[None, GT_LABELS.get(label, label), value] += 1
    return truth, controls, Counter()


def norm(text):
    return " ".join(unicodedata.normalize("NFKC", text).casefold().split())


# a street matches when one holds the other, since detections drop house numbers
def same(label, found, expected):
    found, expected = norm(found), norm(expected)
    return found == expected or (label == "Street" and bool(found) and (found in expected or expected in found))


def word_pattern(value):
    return rf"(?<!\w){re.escape(value)}(?!\w)"


# where each expected value and each control value occurs in a text, as (start, end, label)
def positions(text, values):
    return [(m.start(), m.end(), label) for label, value in values for m in re.finditer(word_pattern(value), text)]


# every detection gets a kind: correct, or by where it falls, control flagged / wrong label / boundary / invented;
# expected values nobody found are added as "missed"
def judge(parts, found, truth, controls):
    open_truth = Counter(truth)
    marks = {part: [] for part in parts}
    for part, occurrences in found.items():
        text = parts[part]
        nearby = [key for key in truth if key[0] in (part, None)]
        expected = positions(text, {(k[1], k[2]) for k in nearby})
        control_spans = positions(text, {(None, value) for value in controls})
        for occurrence in sorted(occurrences, key=lambda o: o["start"]):
            label, start, end = occurrence["label"], occurrence["start"], occurrence["end"]
            key = next((k for k in nearby if open_truth[k] > 0 and k[1] == label and same(label, text[start:end], k[2])),
                       None)
            overlapping = {other for s, e, other in expected if s < end and start < e}
            if key:
                open_truth[key] -= 1
                kind = "correct"
            elif any(s < end and start < e for s, e, _ in control_spans):
                kind = "control flagged"
            elif label in overlapping:
                kind = "boundary"
            elif overlapping:
                kind = "wrong label"
            else:
                kind = "invented"
            marks[part].append({**occurrence, "kind": kind})
    for (part, label, value), count in open_truth.items():
        pattern = word_pattern(value)
        names = [part] if part is not None else [name for name in parts if re.search(pattern, parts[name])]
        for name in names:
            taken = [(m["start"], m["end"]) for m in marks.get(name, []) if "start" in m]
            for match in re.finditer(pattern, parts.get(name, "")):
                if count and not any(s < match.end() and match.start() < e for s, e in taken):
                    marks[name].append({"label": label, "start": match.start(), "end": match.end(), "kind": "missed"})
                    taken.append(match.span())
                    count -= 1
        home = names[0] if names and names[0] in marks else "not in the text"
        marks.setdefault(home, []).extend({"label": label, "value": value, "kind": "missed"} for _ in range(count))
    return marks
""")

code(r"""
# the detector's occurrences per part: the pages go in as one document, the other parts as texts of their own
def detect(detector, parts):
    pages = [name for name in parts if isinstance(name, int)]
    others = [name for name in parts if not isinstance(name, int)]
    starts, position = {}, 0
    for page in pages:
        starts[page] = position
        position += len(parts[page]) + 2
    texts = ["\n\n".join(parts[page] for page in pages)] * bool(pages) + [parts[name] for name in others]
    torch.cuda.synchronize()
    start = time.perf_counter()
    results = detector.detect(texts)
    torch.cuda.synchronize()
    seconds = time.perf_counter() - start
    found = {name: [] for name in parts}
    for occurrence in results[0]["occurrences"] if pages else []:
        page = max((p for p in pages if starts[p] <= occurrence["start"]), key=starts.get)
        end = min(occurrence["end"] - starts[page], len(parts[page]))
        found[page].append({**occurrence, "start": occurrence["start"] - starts[page], "end": end})
    for name, result in zip(others, results[bool(pages):]):
        found[name] = result["occurrences"]
    return found, seconds


# (name, bytes) from an upload, a local path, or the sample of that name in the repository
def read_source(source):
    if isinstance(source, tuple):
        return source
    path = Path(source)
    return path.name, path.read_bytes() if path.exists() else urllib.request.urlopen(SAMPLES + path.name).read()


detectors = {"GLiNER2 fp16": Detector(model=stock, fast=False), "glinert4": Detector(model=fast)}


def evaluate(document, ground_truth=None):
    name, data = read_source(document)
    parts = read_document(name, data)
    truth, controls, other = read_ground_truth(read_source(ground_truth)[1]) if ground_truth else (None, set(), Counter())
    runs = {}
    for setup, detector in detectors.items():
        glinert4.uninstall() if setup == "GLiNER2 fp16" else glinert4.install()
        detector.detect(["Anna Müller of Siemens AG lives at Hauptstraße 5, Berlin."])
        found, seconds = detect(detector, parts)
        marks = judge(parts, found, truth, controls) if truth else {
            part: [{**o, "kind": "detected"} for o in occurrences] for part, occurrences in found.items()}
        runs[setup] = {"seconds": seconds, "found": found, "marks": marks}
    report = {"name": name, "parts": parts, "truth": truth, "controls": controls, "other": other, "runs": runs}
    summary(report)
    return report


def summary(report):
    parts, runs = report["parts"], report["runs"]
    words = sum(len(text.split()) for text in parts.values())
    pages = sum(isinstance(name, int) for name in parts)
    before, after = runs["GLiNER2 fp16"], runs["glinert4"]
    identical = all(sorted((o["label"], o["start"], o["end"]) for o in before["found"][part])
                    == sorted((o["label"], o["start"], o["end"]) for o in after["found"][part]) for part in parts)
    display(Markdown(f"**{report['name']}**: {pages} pages, {words:,} words. GLiNER2 fp16 {before['seconds']:.2f} s, "
                     f"glinert4 {after['seconds']:.2f} s, **{before['seconds'] / after['seconds']:.1f}× faster**, "
                     f"{'identical detections' if identical else 'detections differ'}."))
    if report["truth"] is None:
        return
    rows = []
    for label in dict.fromkeys(key[1] for key in report["truth"]):
        row = [label]
        for run in (before, after):
            kinds = Counter(m["kind"] for marks in run["marks"].values() for m in marks if m["label"] == label)
            tp, fn = kinds["correct"], kinds["missed"]
            fp = sum(kinds.values()) - tp - fn
            row += [f"{tp / max(1, tp + fp):.3f}", f"{tp / max(1, tp + fn):.3f}", fp, fn]
        rows.append(row)
    display(Markdown(table(["Label", "fp16 precision", "fp16 recall", "fp16 FP", "fp16 FN",
                            "glinert4 precision", "glinert4 recall", "glinert4 FP", "glinert4 FN"], rows)))
    kinds = Counter(m["kind"] for marks in after["marks"].values() for m in marks)
    errors = ", ".join(f"{kind} {n}" for kind, n in kinds.items() if kind != "correct") or "none"
    skipped = ", ".join(f"{kind} {n}" for kind, n in report["other"].items())
    display(Markdown(f"glinert4 errors by kind: {errors}. Control values in the ground truth: {len(report['controls'])}."
                     + (f" Not scored by this detector: {skipped}." if skipped else "")))


report = evaluate(DOCUMENT, GROUND_TRUTH)
""")

md("""
### The document, marked up

Every detection in its label's colour. With ground truth, a false positive gets a solid red outline and its kind, and
a missed entity a dashed one where it should have been found. Switch between the two models, jump to a page, or show
only the pages with errors.
""")

code(r"""
ERROR = "#d03b3b"
STYLE = ("<style>"
         ".gt-page{font:14px/1.75 Georgia,'Times New Roman',serif;color:#1d1d1b;background:#fff;border:1px solid #e3e2dd;"
         "border-radius:8px;padding:22px 26px;margin:10px 0;white-space:pre-wrap;max-width:860px;"
         "box-shadow:0 1px 3px #0000000d}"
         ".gt-head{font:600 13px/1.4 system-ui,sans-serif;color:#52514e;margin:18px 0 0;max-width:860px;display:flex;"
         "gap:12px;flex-wrap:wrap;align-items:center}"
         ".gt-head b{color:#1d1d1b;font-size:15px}"
         ".gt-pill{font:500 12px system-ui,sans-serif;border-radius:999px;padding:2px 9px;background:#f0efec;color:#1d1d1b}"
         ".gt-mark{border-radius:3px;padding:1px 2px;color:#1d1d1b}"
         ".gt-tag{font:600 10px system-ui,sans-serif;letter-spacing:.02em;margin-left:3px;vertical-align:1px}"
         ".gt-note{font:12px/1.6 system-ui,sans-serif;color:#52514e;max-width:860px;margin:-2px 0 8px;padding:8px 12px;"
         "border-left:3px solid #d03b3b;background:#d03b3b0d;border-radius:4px}"
         ".gt-legend{font:12px system-ui,sans-serif;color:#52514e;display:flex;gap:16px;flex-wrap:wrap;margin:6px 0 4px}"
         "</style>")


def tag(mark):
    return f'{html.escape(mark["kind"]).upper()} · {html.escape(mark["label"])}'


def mark_html(text, mark, color):
    body = html.escape(text)
    if mark["kind"] in ("correct", "detected"):
        return f'<span class="gt-mark" style="background:{color}22;border-bottom:2px solid {color}">{body}</span>'
    line, fill = ("dashed", "transparent") if mark["kind"] == "missed" else ("solid", f"{color}22")
    return (f'<span class="gt-mark" style="background:{fill};outline:2px {line} {ERROR};outline-offset:1px">{body}'
            f'<span class="gt-tag" style="color:{ERROR}">{tag(mark)}</span></span>')


# one page: its text with every mark, counts on top, and marks that overlap another one listed underneath
def page_html(name, text, marks, palette):
    out, last, aside = [], 0, [m for m in marks if "start" not in m]
    order = {"missed": 2, "correct": 1, "detected": 1}
    for mark in sorted((m for m in marks if "start" in m), key=lambda m: (m["start"], order.get(m["kind"], 0))):
        if mark["start"] < last:
            aside.append(mark)
            continue
        out += [html.escape(text[last:mark["start"]]),
                mark_html(text[mark["start"]:mark["end"]], mark, palette.get(mark["label"], "#888888"))]
        last = mark["end"]
    out.append(html.escape(text[last:]))
    kinds = [m["kind"] for m in marks]
    good = kinds.count("correct") + kinds.count("detected")
    wrong = len(kinds) - good - kinds.count("missed")
    pills = [f'<span class="gt-pill">{good} {"detected" if "detected" in kinds else "correct"}</span>'] * bool(good)
    pills += [f'<span class="gt-pill" style="background:{ERROR}1a;color:{ERROR}">{n} {what}</span>'
              for n, what in ((wrong, "false positive" + "s" * (wrong != 1)), (kinds.count("missed"), "missed")) if n]
    notes = "".join(f'<div>{tag(m)}: "{html.escape(m.get("value") or text[m["start"]:m["end"]])}"</div>' for m in aside)
    title = f"Page {name}" if isinstance(name, int) else str(name).capitalize()
    return (f'<div class="gt-head"><b>{title}</b>{"".join(pills)}</div><div class="gt-page">{"".join(out)}</div>'
            + (f'<div class="gt-note"><b>Not shown inline (overlaps another mark)</b>{notes}</div>' if notes else ""))


def document_html(report, setup, errors_only=False, part="All"):
    marks = report["runs"][setup]["marks"]
    labels = list(dict.fromkeys([*detectors[setup].config.labels, *(m["label"] for ms in marks.values() for m in ms)]))
    palette = dict(zip(labels, PALETTE))
    legend = [f'<span class="gt-mark" style="background:{c}22;border-bottom:2px solid {c}">{html.escape(l)}</span>'
              for l, c in palette.items()]
    if report["truth"] is not None:
        legend += [f'<span class="gt-mark" style="outline:2px solid {ERROR}">false positive</span>',
                   f'<span class="gt-mark" style="outline:2px dashed {ERROR}">missed</span>']
    shown = [name for name in marks if part in ("All", name)
             and (not errors_only or any(m["kind"] not in ("correct", "detected") for m in marks[name]))]
    pages = "".join(page_html(name, report["parts"].get(name, ""), marks[name], palette) for name in shown)
    return STYLE + f'<div class="gt-legend">{"".join(legend)}</div>' + (pages or "<p>Nothing to show.</p>")


setup_choice = widgets.ToggleButtons(options=list(detectors), value="glinert4")
errors_choice = widgets.Checkbox(description="only pages with errors")
part_choice = widgets.Dropdown(options=["All", *report["runs"]["glinert4"]["marks"]], description="Page")
viewer = widgets.Output()


def show(_=None):
    with viewer:
        viewer.clear_output()
        display(HTML(document_html(report, setup_choice.value, errors_choice.value, part_choice.value)))


for choice in (setup_choice, errors_choice, part_choice):
    choice.observe(show, "value")
display(widgets.VBox([widgets.HBox([setup_choice, errors_choice, part_choice]), viewer]))
show()
""")

md("""
The pages where glinert4 made a mistake, as a static copy of the viewer for anyone reading this without running it:
""")

code(r"""
display(HTML(document_html(report, "glinert4", errors_only=report["truth"] is not None)))
""")

md("""
### Upload your own

Pick a document and, optionally, its ground truth, then press Evaluate. The summary appears below and the viewer above
switches to your document.
""")

code(r"""
document_upload = widgets.FileUpload(accept=".docx,.pdf,.txt", description="Document")
truth_upload = widgets.FileUpload(accept=".json", description="Ground truth")
evaluate_button = widgets.Button(description="Evaluate", button_style="primary", icon="search")
upload_output = widgets.Output()


# the first uploaded file as (name, bytes), for ipywidgets 7 and 8
def uploaded(upload):
    files = list(upload.value.values()) if isinstance(upload.value, dict) else list(upload.value)
    if not files:
        return None
    file = files[0]
    return file.get("name") or file["metadata"]["name"], bytes(file["content"])


def on_evaluate(_):
    global report
    with upload_output:
        upload_output.clear_output()
        if uploaded(document_upload) is None:
            print("Choose a document first.")
            return
        report = evaluate(uploaded(document_upload), uploaded(truth_upload))
        part_choice.options = ["All", *report["runs"]["glinert4"]["marks"]]
        show()


evaluate_button.on_click(on_evaluate)
display(widgets.VBox([widgets.HBox([document_upload, truth_upload, evaluate_button]), upload_output]))
""")

md("""
## What made the difference

- **Profile first.** On a T4, attention and the matrix multiplications each took about 45% of the GPU time. cuBLAS was
  already near the hardware limit, so the attention was the part worth rewriting.
- **Write for the GPU you have.** The T4 (sm_75) has fp16 tensor cores but no asynchronous copies, and recent Triton
  releases stop using its tensor cores. The kernel is CUDA C++ with `mma.sync` for this one chip.
- **Exact before clever.** Every optimisation here gives the same predictions up to fp16 rounding; the count head and
  the pipeline are bit-identical to stock GLiNER2.
- **Keep the GPU busy.** Once the kernels were fast, the GPU sat idle while Python prepared the next batch.
  Overlapping the two was the last big step.
""")

notebook = {"cells": cells, "metadata": {"accelerator": "GPU", "colab": {"provenance": [], "gpuType": "T4"},
                                          "kernelspec": {"display_name": "Python 3", "name": "python3"},
                                          "language_info": {"name": "python"}},
            "nbformat": 4, "nbformat_minor": 0}
target = Path(__file__).with_name("glinert4_speed.ipynb")
target.write_text(json.dumps(notebook, indent=1, ensure_ascii=False) + "\n")
print(target)
