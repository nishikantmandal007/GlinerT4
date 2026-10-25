# Reading a document (.docx page by page with tables, tracked deletions, comments, header, footer and properties;
# .pdf text layer; plain text) and scoring detections against a ground-truth manifest per page and occurrence, with
# control values that must not be detected. The same code as section 5 of notebooks/glinert4_speed.ipynb.
import io
import json
import re
import unicodedata
from collections import Counter

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

GT_LABELS = {"Names": "Person", "Person": "Person", "Organization": "Organization", "Street": "Street", "City": "City"}


# a .docx as {page: text} split at explicit page breaks, plus its header, footer, comments and properties
def read_docx(data):
    doc = Document(io.BytesIO(data))
    pages, lines = [], []
    for block in doc.element.body.iterchildren():
        if block.tag == qn("w:p"):
            if block.find(f"{qn('w:pPr')}/{qn('w:pageBreakBefore')}") is not None and lines:
                pages.append("\n".join(lines))
                lines = []
            lines.append(Paragraph(block, doc).text)
            for change in block.iter(qn("w:del")):
                deleted = "".join(t.text or "" for t in change.iter(qn("w:delText")))
                lines.append(f"Deleted by {change.get(qn('w:author'))}: {deleted}")
            if any(br.get(qn("w:type")) == "page" for br in block.iter(qn("w:br"))):
                pages.append("\n".join(lines))
                lines = []
        elif block.tag == qn("w:tbl"):
            lines += ["\t".join(cell.text for cell in row.cells) for row in Table(block, doc).rows]
    if lines:
        pages.append("\n".join(lines))
    parts = dict(enumerate(pages, 1))
    unique = lambda texts: "\n".join(dict.fromkeys(t for t in texts if t.strip()))
    parts["header"] = unique(p.text for s in doc.sections for p in s.header.paragraphs)
    parts["footer"] = unique(p.text for s in doc.sections for p in s.footer.paragraphs)
    parts["comments"] = "\n".join(f"Comment by {c.author}: {c.text}" for c in doc.comments)
    core = doc.core_properties
    parts["properties"] = "\n".join(f"{field.replace('_', ' ').capitalize()}: {getattr(core, field)}"
                                    for field in ("author", "last_modified_by", "title", "subject", "keywords")
                                    if getattr(core, field))
    return {name: text for name, text in parts.items() if text.strip()}


def read_document(name, data):
    if name.lower().endswith(".docx"):
        return read_docx(data)
    if name.lower().endswith(".pdf"):
        from pypdf import PdfReader

        return {page: p.extract_text() or "" for page, p in enumerate(PdfReader(io.BytesIO(data)).pages, 1)}
    return {1: data.decode("utf-8-sig")}


def part_name(location):
    location = location or {}
    part = location.get("part") or "word/document.xml"
    if part.endswith("document.xml"):
        return location.get("page") or 1
    for name in ("header", "footer", "comments"):
        if name in part:
            return name
    return "properties" if part.startswith("docProps/") else part


# Counter{(part, label, value): occurrences}, the control values, and the types this detector does not score
def read_ground_truth(data):
    raw = json.loads(data)
    truth, controls = Counter(), {c["value"] for c in raw.get("controls", [])}
    if "pii" in raw:
        for item in raw["pii"]:
            if item["type"] in GT_LABELS:
                truth[part_name(item.get("location")), GT_LABELS[item["type"]], item["value"]] += 1
        return truth, controls, Counter(item["type"] for item in raw["pii"] if item["type"] not in GT_LABELS)
    for label, values in raw.items():
        for value in values:
            truth[None, GT_LABELS.get(label, label), value] += 1
    return truth, controls, Counter()


def norm(text):
    return " ".join(unicodedata.normalize("NFKC", text).casefold().split())


# a street matches when one holds the other, since detections drop house numbers
def same(label, found, expected):
    found, expected = norm(found), norm(expected)
    return found == expected or (label == "Street" and bool(found) and (found in expected or expected in found))


def word_pattern(value):
    return rf"(?<!\w){re.escape(value)}(?!\w)"


# where each expected value and each control value occurs in a text, as (start, end, label)
def positions(text, values):
    return [(m.start(), m.end(), label) for label, value in values for m in re.finditer(word_pattern(value), text)]


# every detection gets a kind: correct, or by where it falls, control flagged / wrong label / boundary / invented;
# expected values nobody found are added as "missed"
def judge(parts, found, truth, controls):
    open_truth = Counter(truth)
    marks = {part: [] for part in parts}
    for part, occurrences in found.items():
        text = parts[part]
        nearby = [key for key in truth if key[0] in (part, None)]
        expected = positions(text, {(k[1], k[2]) for k in nearby})
        control_spans = positions(text, {(None, value) for value in controls})
        for occurrence in sorted(occurrences, key=lambda o: o["start"]):
            label, start, end = occurrence["label"], occurrence["start"], occurrence["end"]
            key = next((k for k in nearby if open_truth[k] > 0 and k[1] == label and same(label, text[start:end], k[2])),
                       None)
            overlapping = {other for s, e, other in expected if s < end and start < e}
            if key:
                open_truth[key] -= 1
                kind = "correct"
            elif any(s < end and start < e for s, e, _ in control_spans):
                kind = "control flagged"
            elif label in overlapping:
                kind = "boundary"
            elif overlapping:
                kind = "wrong label"
            else:
                kind = "invented"
            marks[part].append({**occurrence, "kind": kind})
    for (part, label, value), count in open_truth.items():
        pattern = word_pattern(value)
        names = [part] if part is not None else [name for name in parts if re.search(pattern, parts[name])]
        for name in names:
            taken = [(m["start"], m["end"]) for m in marks.get(name, []) if "start" in m]
            for match in re.finditer(pattern, parts.get(name, "")):
                if count and not any(s < match.end() and match.start() < e for s, e in taken):
                    marks[name].append({"label": label, "start": match.start(), "end": match.end(), "kind": "missed"})
                    taken.append(match.span())
                    count -= 1
        home = names[0] if names and names[0] in marks else "not in the text"
        marks.setdefault(home, []).extend({"label": label, "value": value, "kind": "missed"} for _ in range(count))
    return marks
